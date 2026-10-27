"""Document text extraction for content search (core/text_extract.py)."""
import io

from app.core.text_extract import extract_text


def test_plain_text():
    out = extract_text(b"hello secret p@ssw0rd here", "notes.txt", "text/plain")
    assert "p@ssw0rd" in out


def test_csv():
    out = extract_text(b"user,pass\nadmin,hunter2\n", "creds.csv", "text/csv")
    assert "hunter2" in out


def test_binary_returns_empty():
    assert extract_text(b"\x00\x01\x02\x03", "blob.bin", "application/octet-stream") == ""


def test_docx():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Confidential finding: SQL injection on /login")
    buf = io.BytesIO()
    doc.save(buf)
    out = extract_text(buf.getvalue(), "report.docx",
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert "SQL injection" in out


def test_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["host", "note"])
    ws.append(["10.0.0.5", "domain admin creds found"])
    buf = io.BytesIO()
    wb.save(buf)
    out = extract_text(buf.getvalue(), "hosts.xlsx",
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "domain admin creds" in out


def test_legacy_xls():
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet1")
    ws.write(0, 0, "host")
    ws.write(0, 1, "note")
    ws.write(1, 0, "10.0.0.9")
    ws.write(1, 1, "legacy xls secret token")
    buf = io.BytesIO()
    wb.save(buf)
    out = extract_text(buf.getvalue(), "old.xls", "application/vnd.ms-excel")
    assert "legacy xls secret token" in out


def test_legacy_doc_graceful_on_garbage():
    # A non-.doc byte blob routed as .doc must not raise — antiword rejects it
    # and we fall back to "".
    assert extract_text(b"not a real doc", "old.doc", "application/msword") == ""


def test_cap_is_enforced():
    from app.core import text_extract

    big = ("word " * 100_000).encode()
    out = extract_text(big, "big.txt", "text/plain")
    assert len(out) <= text_extract.MAX_CHARS
